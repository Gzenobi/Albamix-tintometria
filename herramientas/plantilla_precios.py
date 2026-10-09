"""Genera la planilla Excel de lista de precios de Albamix.

Formato: el mismo del archivo PESO_ESP que exporta la grilla de productos del Albamix original
(Tipo, Código, Nombre, Peso Esp., Precio 1, Capa. 1, Precio 2, Capa. 2, Unidad, Bonif., Rentab.),
con encabezado de vigencia, validaciones y columnas de control de precio por litro.

Uso:
    python -I herramientas/plantilla_precios.py --precios precios.mdb [--peso-esp PESO_ESP] --salida lista-precios.xlsx

- precios.mdb aporta el catálogo actual (productos GVA y bases).
- PESO_ESP (opcional) aporta precios reales; los productos sin precio quedan marcados para completar.
"""
import argparse, csv, datetime, io, os, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from jet1 import parse, table_rows  # noqa: E402

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, AMARILLO_SUAVE, GRIS = '005192', 'FFF5D1', 'F2F4F6'
ENC = ['Tipo', 'Código', 'Nombre', 'Peso Esp.', 'Precio 1', 'Capa. 1', 'Precio 2', 'Capa. 2', 'Unidad', 'Bonif.', 'Rentab.',
       '$ por litro 1', '$ por litro 2']
FILA_ENC = 7


def leer_catalogo(ruta_mdb):
    d = open(ruta_mdb, 'rb').read()
    cols = [('COMP', 'f', 4, 4), ('PE', 'f', 8, 7), ('PRECIO1', 'f', 8, 7), ('CAPA1', 'f', 8, 7), ('PRECIO2', 'f', 8, 7),
            ('CAPA2', 'f', 8, 7), ('NOMBRE', 'v', 60, 10), ('TIPO', 'v', 1, 10), ('UNIDAD', 'v', 1, 10)]
    out = {}
    for _, r in table_rows(d, 18):
        f = parse(r, cols)
        out[f['COMP']] = dict(tipo=f['TIPO'], codigo=f['COMP'], nombre=' '.join(f['NOMBRE'].split()), pe=f['PE'],
                              capa1=f['CAPA1'], capa2=f['CAPA2'], unidad='Kilos' if f['UNIDAD'] == 'K' else 'Litros',
                              precio1=None, precio2=None)
    return out


def leer_peso_esp(ruta):
    filas = csv.DictReader(io.StringIO(open(ruta, encoding='cp1252').read()), delimiter=';')
    return {int(r['Código']): r for r in filas}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--precios', required=True)
    ap.add_argument('--peso-esp')
    ap.add_argument('--salida', default='lista-precios.xlsx')
    a = ap.parse_args()

    prods = leer_catalogo(a.precios)
    if a.peso_esp:
        for cod, r in leer_peso_esp(a.peso_esp).items():
            p = prods.get(cod) or dict(tipo=r['Tipo'], codigo=cod, nombre=r['Nombre'].strip(), pe=float(r['Peso Esp.']),
                                       capa1=float(r['Capa. 1']), capa2=float(r['Capa. 2']), unidad=r['Unidad'])
            p['precio1'] = float(r['Precio 1']) or None
            p['precio2'] = float(r['Precio 2']) or None
            if not p.get('capa2'):
                p['capa2'] = float(r['Capa. 2'])
            prods[cod] = p
    filas = sorted(prods.values(), key=lambda p: ({'B': 0, 'C': 1, 'A': 2}.get(p['tipo'], 3), p['codigo']))

    wb = Workbook()
    ws = wb.active
    ws.title = 'Precios'
    base = Font(name='Arial', size=10)
    entrada = Font(name='Arial', size=10, color='0000FF')
    negrita = Font(name='Arial', size=10, bold=True)
    fino = Side(style='thin', color='B7B9BA')

    ws['A1'] = 'Lista de precios Albamix'
    ws['A1'].font = Font(name='Arial', size=14, bold=True, color=NAVY)
    meta = [('Vigente desde:', datetime.date.today(), 'Fecha desde la que rige esta lista (dd/mm/aaaa).'),
            ('Número de lista:', datetime.date.today().strftime('%Y-%m'), 'Identificador de la lista, por ejemplo 2026-10. Cada lista nueva debe ser posterior a la anterior.'),
            ('Notas:', '', 'Texto libre que verán los distribuidores al actualizar (opcional).')]
    for i, (et, val, ayuda) in enumerate(meta, start=2):
        ws.cell(i, 1, et).font = negrita
        c = ws.cell(i, 2, val)
        c.font = entrada
        c.fill = PatternFill('solid', fgColor=AMARILLO_SUAVE)
        c.comment = Comment(ayuda, 'Albamix')
        if isinstance(val, datetime.date):
            c.number_format = 'dd/mm/yyyy'
    ws.merge_cells('B4:H4')
    ws['A5'] = ('Celdas en azul: datos que se cargan. Columnas grises: control automático de precio por litro (no editar). '
                'Celdas amarillas: precios que faltan, completar antes de distribuir. Ver la hoja Instrucciones.')
    ws['A5'].font = Font(name='Arial', size=9, italic=True, color='868688')
    ws.merge_cells('A5:M5')

    for j, h in enumerate(ENC, start=1):
        c = ws.cell(FILA_ENC, j, h)
        c.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=NAVY if j <= 11 else '868688')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    notas_enc = {
        'Tipo': 'B = Base, C = Colorante, A = Accesorio.',
        'Código': 'Código de producto de 7 dígitos.',
        'Peso Esp.': 'Peso específico en kg/L. Se usa para pasar de ml a gramos en la balanza.',
        'Precio 1': 'Precio del envase chico, sin IVA, en pesos.',
        'Capa. 1': 'Capacidad del envase chico (litros o kilos según Unidad).',
        'Precio 2': 'Precio del envase grande, sin IVA. Dejar en 0 si no existe.',
        'Capa. 2': 'Capacidad del envase grande. Dejar en 0 si no existe.',
        'Unidad': 'Litros o Kilos.',
        'Bonif.': 'Bonificación en % (opcional, 0 si no aplica).',
        'Rentab.': 'Rentabilidad sugerida en % (opcional). El distribuidor puede definir la suya en la app.',
        '$ por litro 1': 'Control: Precio 1 / Capa. 1. Sirve para detectar errores de carga.',
        '$ por litro 2': 'Control: Precio 2 / Capa. 2. Debería ser algo menor que el de la lata chica.',
    }
    for j, h in enumerate(ENC, start=1):
        if h in notas_enc:
            ws.cell(FILA_ENC, j).comment = Comment(notas_enc[h], 'Albamix')

    for i, p in enumerate(filas, start=FILA_ENC + 1):
        capa2 = p['capa2'] or 0
        precio2 = p['precio2'] if p['precio2'] else (None if capa2 > 0 else 0)
        vals = [p['tipo'], p['codigo'], p['nombre'], p['pe'], p['precio1'], p['capa1'], precio2, capa2, p['unidad'], 0, 0]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(i, j, v)
            c.font = entrada
            c.border = Border(bottom=fino)
        ws.cell(i, 4).number_format = '0.000'
        for j in (5, 7):
            ws.cell(i, j).number_format = '#,##0.00'
        for j in (6, 8):
            ws.cell(i, j).number_format = '0.00'
        for j, falta in ((5, p['precio1'] is None), (7, precio2 is None)):
            if falta:
                ws.cell(i, j).fill = PatternFill('solid', fgColor='FFE08A')
                ws.cell(i, j).comment = Comment('Sin precio: completar antes de distribuir la lista.', 'Albamix')
        ws.cell(i, 12, f'=IF(AND(ISNUMBER(E{i}),N(F{i})>0),E{i}/F{i},"")')
        ws.cell(i, 13, f'=IF(AND(ISNUMBER(G{i}),N(G{i})>0,N(H{i})>0),G{i}/H{i},"")')
        for j in (12, 13):
            c = ws.cell(i, j)
            c.font = base
            c.fill = PatternFill('solid', fgColor=GRIS)
            c.number_format = '#,##0.00'
            c.border = Border(bottom=fino)
    ultima = FILA_ENC + len(filas)
    rango = lambda col: f'{col}{FILA_ENC + 1}:{col}{ultima + 500}'

    dv_tipo = DataValidation(type='list', formula1='"B,C,A"', allow_blank=False, showErrorMessage=True,
                             errorTitle='Tipo inválido', error='Usá B (base), C (colorante) o A (accesorio).')
    dv_uni = DataValidation(type='list', formula1='"Litros,Kilos"', showErrorMessage=True,
                            errorTitle='Unidad inválida', error='Usá Litros o Kilos.')
    dv_cod = DataValidation(type='whole', operator='between', formula1='1000000', formula2='9999999', showErrorMessage=True,
                            errorTitle='Código inválido', error='El código de producto tiene 7 dígitos.')
    dv_pos = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', showErrorMessage=True,
                            errorTitle='Valor inválido', error='Ingresá un número mayor o igual a cero.')
    for dv, cols in ((dv_tipo, 'A'), (dv_uni, 'I'), (dv_cod, 'B'), (dv_pos, 'DEFGHJK')):
        ws.add_data_validation(dv)
        for col in cols:
            dv.add(rango(col))

    anchos = [6, 10, 46, 9, 12, 8, 12, 8, 9, 7, 8, 12, 12]
    for j, w in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[FILA_ENC].height = 30
    ws.freeze_panes = ws.cell(FILA_ENC + 1, 4)
    ws.auto_filter.ref = f'A{FILA_ENC}:M{ultima}'
    for row in ws.iter_rows(min_row=1, max_row=6):
        for c in row:
            if c.font and c.font.name != 'Arial':
                c.font = base

    ins = wb.create_sheet('Instrucciones')
    texto = [
        ('Cómo usar esta planilla', True),
        ('1. Completá o corregí los datos en azul de la hoja Precios. Los precios van sin IVA, en pesos.', False),
        ('2. Actualizá "Vigente desde" y "Número de lista" (cada lista nueva tiene que ser posterior a la anterior).', False),
        ('3. Completá los precios marcados en amarillo: son productos que todavía no tienen precio.', False),
        ('4. Revisá las columnas grises de $ por litro: un valor muy distinto al resto suele ser un error de carga.', False),
        ('5. Guardá el archivo y generá el paquete de actualización desde la app (modo administrador).', False),
        ('', False),
        ('Reglas que verifica la app al leer la planilla', True),
        ('- Código de 7 dígitos, sin repetir.', False),
        ('- Tipo B, C o A. Unidad Litros o Kilos. Peso específico mayor que cero.', False),
        ('- Precio 1 y Capa. 1 obligatorios y mayores que cero.', False),
        ('- Precio 2 y Capa. 2: los dos completos o los dos en cero.', False),
        ('- Se pueden agregar productos nuevos al final; un producto que se quita de la planilla no se borra de la app.', False),
        ('', False),
        ('Ejemplo de una fila', True),
        ('B | 4510000 | Esm. Sintetico Industrial Transparente | 0,910 | 20,40 | 3,20 | 93,47 | 16,00 | Litros | 0 | 0', False),
        ('', False),
        ('Formato tomado del archivo PESO_ESP que exporta la grilla de productos del Albamix original.', False),
    ]
    for i, (t, neg) in enumerate(texto, start=1):
        c = ins.cell(i, 1, t)
        c.font = Font(name='Arial', size=12 if neg and i == 1 else 10, bold=neg, color=NAVY if neg else '000000')
    ins.column_dimensions['A'].width = 110

    wb.save(a.salida)
    sin_precio = sum(1 for p in filas if p['precio1'] is None)
    print(f'Planilla creada: {a.salida} ({len(filas)} productos, {sin_precio} sin precio para completar)')


if __name__ == '__main__':
    main()
